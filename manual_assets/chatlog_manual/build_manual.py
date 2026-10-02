from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/teresamarquez/PycharmProjects/ChatLog_v1.5.0")
RAW_SHOTS = Path("/private/tmp/chatlog_manual_assets")
WORK_DIR = ROOT / "manual_assets" / "chatlog_manual"
OUT_DIR = ROOT / "manuals"
DOCX_PATH = OUT_DIR / "Manual_Usuario_ChatLog_v2.0.1.docx"

NAVY = RGBColor(40, 38, 102)
GOLD = RGBColor(255, 186, 40)
TEXT = RGBColor(32, 34, 43)
MUTED = RGBColor(94, 99, 115)
LIGHT_FILL = "F6F4FF"
TIP_FILL = "FFF5D6"
LINE = "D8D7E5"


@dataclass(frozen=True)
class Callout:
    number: int
    x: int
    y: int
    label: str


ANNOTATIONS: dict[str, list[Callout]] = {
    "01_onboarding.png": [
        Callout(1, 34, 83, "Botón para cerrar la guía cuando ya entendiste el recorrido general."),
        Callout(2, 70, 610, "Opción para omitir la guía y entrar de inmediato al trabajo."),
        Callout(3, 88, 684, "Botón para avanzar por los pasos de la guía."),
    ],
    "02_proyectos.png": [
        Callout(1, 56, 134, "Pestaña principal para organizar el trabajo por proyectos."),
        Callout(2, 58, 280, "Ruta de trabajo resumida para comenzar sin perderse."),
        Callout(3, 89, 454, "Acceso directo a la gestión de proyectos."),
        Callout(4, 115, 723, "Botón para ir a la generación de declaraciones."),
    ],
    "03_declaracion.png": [
        Callout(1, 149, 132, "Pestaña para transformar registros en una declaración de uso."),
        Callout(2, 107, 348, "Filtro para elegir qué registros incluir."),
        Callout(3, 129, 561, "Menú para escoger el formato de salida."),
    ],
    "04_herramientas_superior.png": [
        Callout(1, 253, 133, "Pestaña donde se concentran respaldos, importación y mantenimiento."),
        Callout(2, 164, 390, "Selector del alcance de exportación."),
        Callout(3, 156, 489, "Botón para descargar un respaldo completo en JSON."),
        Callout(4, 161, 545, "Botón para exportar registros en CSV."),
    ],
    "05_herramientas_exportacion.png": [
        Callout(1, 110, 222, "Activar recordatorios automáticos de respaldo."),
        Callout(2, 167, 314, "Frecuencia sugerida de recordatorio."),
        Callout(3, 156, 532, "Modo de incorporación de datos al importar."),
        Callout(4, 108, 742, "Selección del archivo JSON a restaurar."),
    ],
    "06_herramientas_revision.png": [
        Callout(1, 131, 215, "Criterios que ayudan a revisar la calidad de los registros."),
        Callout(2, 87, 382, "Guardar los criterios configurados."),
        Callout(3, 207, 511, "Resumen del estado de los registros."),
    ],
    "07_estadisticas.png": [
        Callout(1, 302, 134, "Pestaña para observar patrones de uso."),
        Callout(2, 211, 320, "Indicadores rápidos del total de interacciones, modelos y proyectos."),
        Callout(3, 52, 774, "Zona de filtros para refinar los datos mostrados."),
    ],
    "08_sobre.png": [
        Callout(1, 78, 108, "Panel con contexto académico, versión y enlaces de apoyo."),
        Callout(2, 352, 129, "Botón para cerrar el panel Sobre."),
        Callout(3, 73, 423, "Sección con la versión actual de la extensión."),
    ],
}


def set_page(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_paragraph_border(paragraph, color: str = "282666", size: str = "8") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def style_document(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.font.color.rgb = TEXT

    for style_name, size, color in [
        ("Title", 24, NAVY),
        ("Heading 1", 18, NAVY),
        ("Heading 2", 14, NAVY),
        ("Heading 3", 12, NAVY),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True


def add_title_block(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("Manual de usuario de ChatLog")
    run.font.name = "Arial"
    run.font.size = Pt(25)
    run.font.bold = True
    run.font.color.rgb = NAVY

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(18)
    run = sub.add_run("Version 2.0.1 | Guia para profesores universitarios y estudiantes")
    run.font.name = "Arial"
    run.font.size = Pt(12)
    run.font.color.rgb = MUTED

    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.paragraph_format.space_after = Pt(18)
    run = info.add_run(
        "Documento pedagogico de consulta rapida para registrar, organizar y declarar el uso academico de herramientas de inteligencia artificial."
    )
    run.font.name = "Arial"
    run.font.size = Pt(11)

    rule = doc.add_paragraph()
    set_paragraph_border(rule)


def add_note_box(doc: Document, title: str, lines: Iterable[str], fill: str = LIGHT_FILL) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    r.bold = True
    r.font.name = "Arial"
    r.font.color.rgb = NAVY
    for line in lines:
        p = cell.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(line)
        run.font.name = "Arial"
        run.font.size = Pt(10.5)
    doc.add_paragraph()


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    heading = doc.add_heading(text, level=level)
    heading.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    heading.paragraph_format.space_after = Pt(6)


def add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15


def add_numbered_steps(doc: Document, steps: Iterable[str]) -> None:
    for step in steps:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        p.add_run(step)


def add_bullets(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        p.add_run(item)


def annotate_image(raw_name: str) -> Path:
    source = RAW_SHOTS / raw_name
    target = WORK_DIR / f"annotated_{raw_name}"
    if not source.exists() and target.exists():
        return target
    image = Image.open(source).convert("RGB")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial Bold.ttf", 24)

    if raw_name != "01_onboarding.png":
        draw.rectangle((0, 72, 250, 106), fill="white")

    for callout in ANNOTATIONS.get(raw_name, []):
        draw.ellipse(
            (callout.x - 18, callout.y - 18, callout.x + 18, callout.y + 18),
            fill=(255, 186, 40),
            outline=(40, 38, 102),
            width=3,
        )
        text = str(callout.number)
        bbox = draw.textbbox((0, 0), text, font=font)
        text_w = bbox[2] - bbox[0]
        text_h = bbox[3] - bbox[1]
        draw.text(
            (callout.x - text_w / 2, callout.y - text_h / 2 - 1),
            text,
            fill=(40, 38, 102),
            font=font,
        )

    image.save(target)
    return target


def add_annotated_capture(doc: Document, title: str, image_name: str) -> None:
    add_heading(doc, title, level=2)
    image_path = annotate_image(image_name)
    doc.add_picture(str(image_path), width=Inches(3.55))
    figure = doc.paragraphs[-1]
    figure.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure.paragraph_format.space_after = Pt(4)
    for callout in ANNOTATIONS[image_name]:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(callout.label)
        run.font.name = "Arial"
        run.font.size = Pt(10.5)
    doc.add_paragraph()


def add_example_table(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    headers = ["Perfil", "Uso sugerido de ChatLog", "Ejemplo de registro"]
    for idx, header in enumerate(headers):
        cell = table.cell(0, idx)
        set_cell_shading(cell, "E8EEF5")
        run = cell.paragraphs[0].add_run(header)
        run.bold = True
        run.font.name = "Arial"
    rows = [
        (
            "Profesorado",
            "Documentar apoyo en planeacion, materiales y evaluacion.",
            "Uso de Gemini para generar una primera lista de preguntas para una practica de aula.",
        ),
        (
            "Estudiantes",
            "Registrar consultas de lectura, resumen y apoyo para redactar borradores.",
            "Uso de ChatGPT para comparar dos definiciones teoricas y preparar un esquema de ensayo.",
        ),
        (
            "Investigacion",
            "Guardar trazabilidad de prompts, verificacion y etapa del trabajo.",
            "Uso de Claude para sintetizar temas recurrentes en notas de entrevista exploratoria.",
        ),
    ]
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
            for p in cells[idx].paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(10.5)
    doc.add_paragraph()


def build_manual() -> None:
    WORK_DIR.mkdir(parents=True, exist_ok=True)
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    doc = Document()
    set_page(doc.sections[0])
    style_document(doc)
    add_title_block(doc)

    add_note_box(
        doc,
        "Guia rapida de 1 pagina",
        [
            "1. Abre ChatLog y crea o selecciona un proyecto.",
            "2. Registra cada interaccion importante con la herramienta de IA.",
            "3. Describe con tus palabras para que usaste la herramienta y como verificaste el resultado.",
            "4. Cuando termines, genera una declaracion de uso y exporta un respaldo en JSON o CSV.",
            "5. Si trabajas en equipo, usa nombres de proyecto claros y exporta un respaldo por proyecto.",
        ],
        fill=TIP_FILL,
    )

    add_heading(doc, "1. Introduccion", level=1)
    add_body(
        doc,
        "ChatLog es una extension de Chrome disenada para documentar de manera clara y trazable el uso de modelos de lenguaje en actividades academicas. Su objetivo es ayudar a profesores y estudiantes a registrar como utilizaron estas herramientas, con que finalidad y bajo que criterios de revision humana."
    )
    add_body(
        doc,
        "Este manual esta pensado para personas no tecnicas. Por eso presenta instrucciones paso a paso, capturas comentadas, ejemplos concretos de uso universitario y recomendaciones practicas para evitar errores comunes."
    )

    add_heading(doc, "2. Requisitos previos", level=1)
    add_bullets(
        doc,
        [
            "Contar con Google Chrome actualizado.",
            "Tener instalada la extension ChatLog desde la Chrome Web Store o desde la fuente institucional autorizada.",
            "Disponer de acceso a internet si se van a abrir las conversaciones originales en plataformas externas.",
            "Tener claro el contexto academico del trabajo: curso, proyecto, tesis, articulo, actividad o practica.",
        ],
    )

    add_heading(doc, "3. Acceso al sistema", level=1)
    add_numbered_steps(
        doc,
        [
            "Abre Google Chrome.",
            "Localiza el icono de ChatLog en la barra de extensiones.",
            "Haz clic en el icono para abrir el panel lateral.",
            "Si es tu primera vez, revisa la guia de inicio y despues cierrala o pulsa Siguiente hasta terminarla.",
        ],
    )
    add_note_box(
        doc,
        "Recomendacion",
        [
            "Si no ves el icono, revisa el menu de extensiones de Chrome y fija ChatLog para tenerlo siempre a la mano.",
        ],
    )

    add_heading(doc, "4. Descripcion de la interfaz", level=1)
    add_body(
        doc,
        "La interfaz se organiza como un panel lateral con cuatro pestañas principales: Proyectos, Declaracion de uso, Herramientas y Estadisticas. Ademas, ofrece una guia de inicio y un panel Sobre para conocer el contexto del proyecto."
    )
    add_annotated_capture(doc, "4.1 Guia de inicio", "01_onboarding.png")
    add_annotated_capture(doc, "4.2 Pestaña Proyectos", "02_proyectos.png")
    add_annotated_capture(doc, "4.3 Pestaña Declaracion de uso", "03_declaracion.png")
    add_annotated_capture(doc, "4.4 Pestaña Herramientas", "04_herramientas_superior.png")
    add_annotated_capture(doc, "4.5 Importacion y respaldo", "05_herramientas_exportacion.png")
    add_annotated_capture(doc, "4.6 Revision de registros", "06_herramientas_revision.png")
    add_annotated_capture(doc, "4.7 Pestaña Estadisticas", "07_estadisticas.png")
    add_annotated_capture(doc, "4.8 Panel Sobre", "08_sobre.png")

    doc.add_section(WD_SECTION.NEW_PAGE)
    set_page(doc.sections[-1])

    add_heading(doc, "5. Procedimientos paso a paso", level=1)
    add_heading(doc, "5.1 Crear un proyecto", level=2)
    add_numbered_steps(
        doc,
        [
            "En la pestaña Proyectos, escribe un nombre claro en el campo de nuevo proyecto.",
            "Usa nombres que permitan identificar facilmente la actividad: por ejemplo, Historia oral 2026 o Seminario IA y etica.",
            "Haz clic en Crear proyecto.",
            "Verifica que el proyecto quede seleccionado antes de registrar una interaccion.",
        ],
    )

    add_heading(doc, "5.2 Registrar una interaccion", level=2)
    add_numbered_steps(
        doc,
        [
            "Completa el nombre de la interaccion con una descripcion breve y util.",
            "Selecciona la finalidad de la interaccion.",
            "Escribe la empresa proveedora, el nombre del modelo y, si aplica, la version.",
            "Pega la liga de la conversacion si la tienes disponible.",
            "Copia el prompt principal o resume la instruccion utilizada.",
            "Explica la etapa del trabajo y como realizaste la verificacion humana.",
            "Espera el indicador Guardado o Borrador guardado. No necesitas pulsar Guardar registro.",
            "Pulsa Nuevo registro para iniciar otra interaccion; la anterior se conserva.",
            "Al volver a abrir el panel se recupera el formulario activo, incluso si estaba incompleto.",
        ],
    )

    add_heading(doc, "5.3 Generar una declaracion de uso", level=2)
    add_numbered_steps(
        doc,
        [
            "Abre la pestaña Declaracion de uso.",
            "Filtra por proyecto si deseas trabajar solo con una actividad concreta.",
            "Selecciona los registros que deban aparecer en la declaracion.",
            "Elige el tipo de trabajo o el formato de declaracion.",
            "Genera la declaracion y revisa el texto antes de copiarlo o descargarlo.",
        ],
    )

    add_heading(doc, "5.4 Exportar un respaldo", level=2)
    add_numbered_steps(
        doc,
        [
            "Abre la pestaña Herramientas.",
            "En Respaldo y transferencia, elige el alcance de la exportacion.",
            "Selecciona JSON si deseas conservar la estructura completa de ChatLog.",
            "Selecciona CSV si deseas abrir los registros en una hoja de calculo.",
            "Guarda el archivo en una carpeta facil de localizar.",
        ],
    )

    add_heading(doc, "5.5 Importar un respaldo", level=2)
    add_numbered_steps(
        doc,
        [
            "En la pestaña Herramientas, localiza el apartado Importar datos.",
            "Decide si deseas fusionar la informacion o reemplazar la actual.",
            "Selecciona el archivo JSON o CSV correspondiente.",
            "Confirma la importacion y revisa el reporte final.",
        ],
    )

    add_heading(doc, "6. Ejemplos concretos de uso", level=1)
    add_body(
        doc,
        "Los siguientes ejemplos muestran como puede emplearse ChatLog en actividades universitarias frecuentes."
    )
    add_example_table(doc)
    add_note_box(
        doc,
        "Ejemplos breves",
        [
            "Docencia: preparar una lista inicial de preguntas de discusion para clase y dejar constancia de que despues fueron revisadas y ajustadas por el profesor.",
            "Estudio: registrar una consulta para resumir un articulo, dejando claro que el texto final fue reescrito y verificado por el estudiante.",
            "Investigacion: documentar prompts usados para explorar categorias analiticas preliminares antes de pasar a una codificacion humana mas cuidadosa.",
        ],
    )

    add_heading(doc, "7. Advertencias sobre errores comunes", level=1)
    add_bullets(
        doc,
        [
            "Guardar registros sin proyecto: dificulta encontrarlos despues.",
            "Escribir nombres vagos como Prueba 1 o Chat IA: hace mas dificil recordar el proposito real de la interaccion.",
            "Omitir la verificacion humana: debilita la trazabilidad academica del trabajo.",
            "Confiar solo en CSV como respaldo maestro: para restaurar la estructura completa conviene conservar tambien un JSON.",
            "No exportar respaldos periodicamente: aumenta el riesgo de perdida de informacion si cambias de equipo o de perfil de navegador.",
        ],
    )

    add_heading(doc, "8. Solucion de problemas", level=1)
    troubleshooting = [
        ("No veo ChatLog en Chrome.", "Revisa el menu de extensiones y verifica si la extension esta instalada y fijada."),
        ("No puedo generar la declaracion.", "Confirma que existan registros guardados y seleccionados para la declaracion."),
        ("El archivo importado no funciona.", "Comprueba que el archivo sea un respaldo JSON o CSV compatible con ChatLog."),
        ("No encuentro un registro.", "Revisa que estes en el proyecto correcto o usa los filtros de Herramientas."),
        ("Quiero mover mis registros a otro equipo.", "Usa Exportar JSON para conservar la estructura completa."),
    ]
    for problem, answer in troubleshooting:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.add_run(f"{problem} ").bold = True
        p.add_run(answer)
    doc.add_paragraph()

    add_heading(doc, "9. Preguntas frecuentes", level=1)
    faq = [
        ("Que diferencia hay entre JSON y CSV?", "JSON conserva la estructura interna de ChatLog; CSV sirve mejor para abrir y revisar registros en una hoja de calculo."),
        ("Debo registrar todas mis interacciones con IA?", "Conviene registrar al menos las que influyeron en decisiones, redaccion, analisis o productos academicos relevantes."),
        ("Puedo usar ChatLog aunque trabaje solo?", "Si. De hecho, ayuda a organizar mejor la memoria del proceso de trabajo."),
        ("Para que sirve la calificacion por estrellas?", "Permite valorar rapidamente que tan util fue una interaccion para recuperarla o compararla despues."),
        ("Las estadisticas reemplazan la revision cualitativa?", "No. Son una ayuda para observar patrones, no un sustituto de la interpretacion academica."),
    ]
    for question, answer in faq:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.add_run(f"{question} ").bold = True
        p.add_run(answer)
    doc.add_paragraph()

    add_heading(doc, "10. Glosario breve", level=1)
    glossary = [
        ("Prompt", "Instruccion o pregunta que se introduce en una herramienta de inteligencia artificial."),
        ("Respaldo JSON", "Archivo que conserva proyectos, registros y estructura interna para poder restaurar la informacion en ChatLog."),
        ("CSV", "Archivo de valores separados por comas, util para abrir registros en hojas de calculo."),
        ("Declaracion de uso", "Texto que resume y transparenta como se utilizo una herramienta de IA dentro de un trabajo academico."),
        ("Verificacion humana", "Revision hecha por una persona para confirmar, corregir o contextualizar la informacion generada por la IA."),
        ("Proyecto", "Conjunto de registros agrupados alrededor de una actividad, curso, investigacion o entrega especifica."),
    ]
    for term, definition in glossary:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.add_run(f"{term}: ").bold = True
        p.add_run(definition)
    doc.add_paragraph()

    add_heading(doc, "11. Soporte y retroalimentacion", level=1)
    add_body(
        doc,
        "ChatLog es un proyecto academico de acceso abierto. Si deseas reportar mejoras, compartir observaciones de uso o sugerir nuevas funciones, conviene canalizar la retroalimentacion por las vias institucionales del proyecto y por el repositorio publico de ChatLog."
    )
    add_bullets(
        doc,
        [
            "Repositorio del proyecto: https://github.com/tmarquez-mx/ChatLog",
            "Responsable academica indicada en la extension: Dra. Teresa Márquez. Contacto: teresa.marquez@ibero.mx.",
            "Politica de privacidad: disponible en el panel Sobre y prevista para https://tmarquez-mx.github.io/ChatLog/privacidad/.",
            "Antes de reportar un problema, anota que estabas intentando hacer, en que pestaña estabas y si el problema ocurrio al exportar, importar o generar una declaracion.",
        ],
    )

    doc.save(DOCX_PATH)


if __name__ == "__main__":
    build_manual()
